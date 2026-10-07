"""Generate synthetic commercial proposals and audit extracted form fields.

Run from the repository root: python -m scripts.quality_audit [API_URL]
This live audit sends synthetic documents to OpenAI and uses API quota.
All generated names, addresses and document values are fictional. Use --env-file
to load configuration; missing API configuration stops before extraction.
"""

import argparse
import json
import os
from pathlib import Path

import fitz
import httpx
import openpyxl
import xlwt
from docx import Document
from PIL import Image
from dotenv import load_dotenv


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "examples" / "mock_kp"
FIELDS = [
    ("Клиент", "ТОО Северный Вектор", "client"),
    ("Контактное лицо", "Айжан М.", "clientContact"),
    ("Поставщик", "ТОО Орион Снаб", "supplier"),
    ("Действительно до", "31.10.2026", "validUntil"),
    ("Валюта", "KZT", "currency"),
    ("НДС", "включён 12%", "vat"),
    ("Условия оплаты", "50% предоплата", "paymentTerms"),
    ("Срок поставки", "7 рабочих дней", "deliveryTerms"),
]
TITLE = "Коммерческое предложение № КП-107 от 30.09.2026"
HEADER = ["№", "Наименование", "Ед. изм.", "Количество", "Цена", "Сумма"]
ROWS = [
    [1, "Кабель медный", "м", 2, 1250, 2500],
    [2, "Кронштейн стальной", "шт.", 3, 400, 1200],
]
EXPECTED = {
    "title": TITLE,
    "client": "ТОО Северный Вектор",
    "clientContact": "Айжан М.",
    "supplier": "ТОО Орион Снаб",
    "validUntil": "31.10.2026",
    "currency": "KZT",
    "vat": "включён 12%",
    "paymentTerms": "50% предоплата",
    "deliveryTerms": "7 рабочих дней",
    "documentNumber": "КП-107",
    "documentDate": "30.09.2026",
    "documentTotal": 3700,
    "items": [
        {"name": "Кабель медный", "unit": "м", "quantity": 2, "unitPrice": 1250, "lineTotal": 2500},
        {"name": "Кронштейн стальной", "unit": "шт.", "quantity": 3, "unitPrice": 400, "lineTotal": 1200},
    ],
}


def font_path():
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    return next(path for path in candidates if path.exists())


def generate():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    metadata_lines = [TITLE, *(f"{label}: {value}" for label, value, _ in FIELDS)]
    table = [HEADER, *ROWS]
    plain = "\n".join(metadata_lines + ["\t".join(map(str, row)) for row in table] + ["Итого: 3 700"])
    (OUTPUT / "offer.txt").write_text(plain, encoding="utf-8")
    (OUTPUT / "offer.csv").write_text(
        "\n".join([TITLE] + [f'"{label}";"{value}"' for label, value, _ in FIELDS]
                  + [";".join(map(str, row)) for row in table] + ["Итого;3700"]),
        encoding="utf-8-sig",
    )
    (OUTPUT / "offer.tsv").write_text(plain, encoding="utf-16")
    (OUTPUT / "offer.json").write_text(
        json.dumps({key: value for _, value, key in FIELDS} | {
            "title": TITLE, "documentNumber": "КП-107", "documentDate": "30.09.2026",
            "documentTotal": 3700, "items": EXPECTED["items"],
        }, ensure_ascii=False, indent=2), encoding="utf-8",
    )

    doc = Document()
    doc.add_paragraph(TITLE)
    details = doc.add_table(rows=0, cols=2)
    for label, value, _ in FIELDS:
        cells = details.add_row().cells
        cells[0].text, cells[1].text = label, value
    items = doc.add_table(rows=0, cols=len(HEADER))
    for values in table:
        for cell, value in zip(items.add_row().cells, values):
            cell.text = str(value)
    doc.add_paragraph("Итого: 3 700")
    doc.save(OUTPUT / "offer.docx")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "КП"
    sheet.append([TITLE])
    for label, value, _ in FIELDS:
        sheet.append([label, value])
    for row in table:
        sheet.append(row)
    sheet.append(["Итого", 3700])
    workbook.save(OUTPUT / "offer.xlsx")

    legacy = xlwt.Workbook()
    sheet = legacy.add_sheet("КП")
    records = [[TITLE], *([label, value] for label, value, _ in FIELDS), *table, ["Итого", 3700]]
    for row_no, values in enumerate(records):
        for col_no, value in enumerate(values):
            sheet.write(row_no, col_no, value)
    legacy.save(str(OUTPUT / "offer.xls"))

    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_font(fontname="fixture", fontfile=str(font_path()))
    for index, line in enumerate(metadata_lines):
        page.insert_text((32, 40 + 24 * index), line, fontname="fixture", fontsize=10)
    x_positions = [32, 60, 255, 325, 405, 475]
    for row_index, values in enumerate(table):
        for x, value in zip(x_positions, values):
            page.insert_text((x, 300 + 25 * row_index), str(value), fontname="fixture", fontsize=10)
    page.insert_text((32, 390), "Итого: 3 700", fontname="fixture", fontsize=10)
    pdf.save(OUTPUT / "offer.pdf")
    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    image.save(OUTPUT / "offer.png")
    image.save(OUTPUT / "offer.jpg", quality=95)
    image.save(OUTPUT / "offer.webp", quality=95)
    scanned = fitz.open()
    scanned_page = scanned.new_page(width=595, height=842)
    scanned_page.insert_image(scanned_page.rect, stream=(OUTPUT / "offer.jpg").read_bytes())
    scanned.save(OUTPUT / "offer-scan.pdf")
    scanned.close()
    image.close()
    pdf.close()
    (OUTPUT / "expected.json").write_text(json.dumps(EXPECTED, ensure_ascii=False, indent=2), encoding="utf-8")


def expected_values():
    values = {key: value for key, value in EXPECTED.items() if key != "items"}
    for index, item in enumerate(EXPECTED["items"]):
        values.update({f"items.{index}.{key}": value for key, value in item.items()})
    return values


def actual_value(proposal, key):
    if not key.startswith("items."):
        return proposal.get(key)
    _, index, field = key.split(".")
    items = proposal.get("items", [])
    return items[int(index)].get(field) if int(index) < len(items) else None


def ensure_configuration(api_url=None):
    if api_url:
        try:
            with httpx.Client(timeout=10, follow_redirects=False) as client:
                response = client.get(f"{api_url.rstrip('/')}/api/health")
                response.raise_for_status()
                health = response.json()
        except (httpx.HTTPError, ValueError):
            raise SystemExit('Не удалось проверить /api/health. Проверьте URL и запущенный backend.') from None
        if health.get('provider') != 'openai' or health.get('configured') is not True:
            raise SystemExit('Backend не настроен для OpenAI. Укажите OPENAI_API_KEY в .env сервера и перезапустите backend.')
    elif not os.getenv('OPENAI_API_KEY', '').strip():
        raise SystemExit('OpenAI API не настроен. Укажите OPENAI_API_KEY в .env или окружении; используйте --env-file PATH.')


def valid_evidence(entry, cells):
    if isinstance(entry, list):
        return bool(entry) and all(valid_evidence(value, cells) for value in entry)
    if not isinstance(entry, dict):
        return False
    if entry.get('method') == 'calculated':
        sources = entry.get('sources', [])
        return bool(sources) and all(valid_evidence(value, cells) for value in sources)
    identifier = entry.get('sourceId', entry.get('id'))
    return identifier in cells and entry.get('excerpt') == cells[identifier]['text']


def audit(api_url=None):
    ensure_configuration(api_url)
    from backend.canonical import DEADLINE
    from backend.pipeline import process_document

    expected = expected_values()
    report = {}
    for path in sorted(OUTPUT.glob("offer.*")) + [OUTPUT / "offer-scan.pdf"]:
        name = path.name
        try:
            if api_url:
                with httpx.Client(timeout=DEADLINE + 15, follow_redirects=False) as client:
                    response = client.post(f"{api_url.rstrip('/')}/api/extract", files={"file": (name, path.read_bytes())})
                    response.raise_for_status()
                    result = response.json()
            else:
                result = process_document(path.read_bytes(), name)
            proposal = result["proposal"]
            mismatches = {}
            missing_evidence = []
            metadata = result['metadata']
            cells = {cell['id']: cell for cell in metadata['sourceCells']}
            for field, value in expected.items():
                actual = actual_value(proposal, field)
                if actual != value:
                    mismatches[field] = {"expected": value, "actual": actual}
                if not valid_evidence(metadata['fieldEvidence'].get(field), cells):
                    missing_evidence.append(field)
            verification = metadata['verification']
            reviewed = (verification.get('mode') == 'model' and verification.get('reviewCompleted') is True
                        and verification.get('coverageComplete') is True
                        and (not verification.get('visionRequired') or verification.get('visionComplete') is True))
            report[name] = {
                "status": "mismatch" if mismatches or missing_evidence else "pass" if reviewed else "partial",
                "fields_correct": len(expected) - len(mismatches),
                "fields_total": len(expected),
                "mismatches": mismatches,
                "missing_evidence": missing_evidence,
                "warnings": metadata["warnings"],
                "provider": metadata.get('provider'),
                "outcome": metadata['outcome']['state'],
                "coverage_complete": verification.get('coverageComplete', False),
                "review_completed": verification.get('reviewCompleted', False),
                "vision_pages": metadata.get('routing', {}).get('processed_pages', []),
            }
        except httpx.HTTPStatusError as error:
            report[name] = {"status": "error", "error": f"Backend HTTP {error.response.status_code}"}
        except Exception as error:
            report[name] = {"status": "error", "error": f"{type(error).__name__}: {error}"}
    report_name = "audit-http.json" if api_url else "audit.json"
    (OUTPUT / report_name).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    for name, result in report.items():
        print(f"{name}: {result['status']} ({result.get('fields_correct', 0)}/{len(expected)})")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('api_url', nargs='?', help='Optional running backend URL; its /api/health must report OpenAI configured')
    parser.add_argument('--env-file', type=Path, default=Path('.env'))
    args = parser.parse_args()
    load_dotenv(args.env_file, override=False)
    ensure_configuration(args.api_url)
    print('LIVE OPENAI AUDIT: synthetic documents are sent to the API; usage is billed to the configured project.')
    generate()
    results = audit(args.api_url)
    if any(result["status"] != 'pass' for result in results.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
