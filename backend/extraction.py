"""Native format adapters, adaptive visual routing and source-preserving fusion."""
import csv, json, os, re, time, zipfile
from io import BytesIO, StringIO
from pathlib import Path
from xml.etree import ElementTree
import fitz
from docx import Document
from .canonical import (
    CanonicalCell, CanonicalDocument, Cell, Source, DocumentError,
    NativeWord, DEADLINE, MAX_CHARS, MAX_CELLS, union_boxes,
)

IMAGE_FORMATS = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff'}
FORMATS = {'.pdf', '.docx', '.xlsx', '.xls', '.csv', '.tsv', '.txt', '.json'} | IMAGE_FORMATS
MAX_BYTES = int(os.getenv('MAX_FILE_SIZE_MB', '25')) * 1024**2
MAX_PAGES = int(os.getenv('MAX_PAGES', '50'))
MAX_ROWS = int(os.getenv('MAX_ROWS', '10000'))
MAX_SHEETS = int(os.getenv('MAX_SHEETS', '30'))

def decode(content):
    if content.startswith((b'\xff\xfe', b'\xfe\xff')): return content.decode('utf-16')
    for encoding in ('utf-8-sig', 'cp1251'):
        try: return content.decode(encoding)
        except UnicodeDecodeError: pass
    raise DocumentError('Не удалось определить кодировку: используйте UTF-8, UTF-16 или Windows-1251')

def check_zip(content):
    with zipfile.ZipFile(BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(e.file_size for e in entries) > 80 * 1024**2:
            raise DocumentError('Слишком большой распакованный документ (лимит 80 МБ)')

def add_lines(source, text, block, kind='text', split_blocks=False, **location):
    current, index = None, 0
    for row, line in enumerate(text.splitlines(), 1):
        values = re.split(r'\t|\s{2,}|\|', line)
        row_kind = 'table' if split_blocks and len(values) > 1 else kind
        if split_blocks and row_kind != current: index += 1; current = row_kind
        row_block = f'{block}-{row_kind}-{index}' if split_blocks else block
        for col, value in enumerate(values, 1): source.add(value, row_block, row, col, kind=row_kind, **location)

def _pdf_display_box(page, box):
    """PyMuPDF text/table coordinates are unrotated; pixmaps apply PDF rotation."""
    return list(fitz.Rect(box) * page.rotation_matrix) if box else None


def _pdf_table(source, table, page_no, index, page):
    block = f'page-{page_no}-table-{index}'
    values = table.extract()
    boxes = [box for row in table.rows for box in row.cells if box]
    x_edges = sorted({round(box[x], 2) for box in boxes for x in (0, 2)})
    y_edges = sorted({round(box[y], 2) for box in boxes for y in (1, 3)})
    anchors = {}
    external = bool(table.header.external and any(table.header.names))
    if external:
        for col, name in enumerate(table.header.names, 1):
            box = table.header.cells[col - 1] if col <= len(table.header.cells) else None
            source.add(name, block, 1, col, page=page_no, kind='table',
                       bbox=_pdf_display_box(page, box), source_method='pdf-native')
    for row_no, (values_row, native_row) in enumerate(zip(values, table.rows), 1 + int(external)):
        for col, value in enumerate(values_row, 1):
            box = native_row.cells[col - 1] if col <= len(native_row.cells) else None
            key = tuple(round(float(v), 2) for v in box) if box else None
            anchor = anchors.get(key) if key else None
            item = source.add('' if anchor else value, block, row_no, col,
                              page=page_no, kind='table', bbox=_pdf_display_box(page, box),
                              source_method='pdf-native', merged_into=anchor.id if anchor else None)
            if box and not anchor:
                anchors[key] = item
                item.colspan = max(1, sum(box[0] - .1 <= x < box[2] - .1 for x in x_edges))
                item.rowspan = max(1, sum(box[1] - .1 <= y < box[3] - .1 for y in y_edges))
    canonical = source._table_index.get(block)
    if canonical and external:
        canonical.header_rows = [1]
    return fitz.Rect(union_boxes([table.bbox, table.header.bbox if external else None]))


def _pdf_native_page(source, page, page_no):
    source.ensure_page(page_no, page.rect.width, page.rect.height, page.rotation)
    words = sorted(page.get_text('words', sort=True), key=lambda w: (w[5], w[6], w[7]))
    source.native_words.extend(NativeWord(f'w{page_no}-{index}', str(w[4]), page_no,
                                        _pdf_display_box(page, w[:4]), int(w[5]), int(w[6]), index)
                               for index, w in enumerate(words))
    rotation = page.rotation
    try:
        # PyMuPDF 1.25 find_tables internally rotates the page content when
        # rotation is set. Detect on the logical unrotated grid instead, then
        # apply the same display matrix as words/cells exactly once.
        if rotation:
            page.set_rotation(0)
        tables = page.find_tables().tables
    except (ValueError, RuntimeError):
        tables = []
        source.warnings.append(f'Страница {page_no}: native PDF parser не восстановил таблицы')
    finally:
        if rotation:
            page.set_rotation(rotation)
    boxes = [_pdf_table(source, table, page_no, index, page) for index, table in enumerate(tables)]
    rows = []
    for word in sorted(words, key=lambda w: (round(w[1] / 3), w[0])):
        if any(box.contains(fitz.Rect(word[:4])) for box in boxes):
            continue
        if not rows or abs(rows[-1][0][1] - word[1]) > 3:
            rows.append([])
        rows[-1].append(word)
    chunk_rows = []
    for words_row in rows:
        chunks = []
        for word in sorted(words_row, key=lambda w: w[0]):
            if not chunks or word[0] - chunks[-1][-1][2] > 12:
                chunks.append([])
            chunks[-1].append(word)
        chunk_rows.append(chunks)
    row_widths = [len(chunks) for chunks in chunk_rows]
    inferred_rows = set()
    if not tables:
        for index in range(len(chunk_rows) - 1):
            first, second = chunk_rows[index:index + 2]
            if len(first) >= 3 and len(first) == len(second) and all(
                    abs(a[0][0] - b[0][0]) < 5 for a, b in zip(first, second)):
                inferred_rows.update((index, index + 1))
    previous_kind, inferred_segment = None, 0
    for row_no, (words_row, chunks) in enumerate(zip(rows, chunk_rows), 1):
        segment = sum(box.y1 <= min(w[1] for w in words_row) for box in boxes)
        text_block = f'page-{page_no}-text-{segment}' if boxes else f'page-{page_no}'
        kind = 'text'
        if inferred_rows:
            kind = 'table' if row_no - 1 in inferred_rows else 'text'
            if kind != previous_kind:
                inferred_segment += 1
                previous_kind = kind
            text_block = f'page-{page_no}-{kind}-{inferred_segment}'
        for col, chunk in enumerate(chunks, 1):
            source.add(' '.join(w[4] for w in chunk), text_block, row_no, col,
                       page=page_no, bbox=_pdf_display_box(page, union_boxes([w[:4] for w in chunk])),
                       source_method='pdf-native', kind=kind)
    # Native extraction retains all text. Reading order follows page geometry,
    # including tables that were parsed before surrounding paragraphs.
    page_blocks = [b for b in source.blocks if b.page == page_no]
    def native_block_position(block):
        box = fitz.Rect(block.bbox) * page.derotation_matrix if block.bbox else fitz.Rect()
        return box.y0, box.x0
    for order, block in enumerate(sorted(page_blocks, key=native_block_position)):
        block.reading_order = order
        for cell in source.cells:
            if cell.page == page_no and cell.block == block.id:
                cell.reading_order = order
    text = page.get_text('text')
    reasons = []
    if sum(ch.isalnum() for ch in text) < 20:
        reasons.append('missing_or_sparse_native_text')
    if '\ufffd' in text or text.count('\x00') > 2:
        reasons.append('damaged_native_text')
    page_tables = [t for t in source.table_objects if t.page == page_no]
    if len(page_tables) > 1:
        reasons.append('multiple_tables')
    if any(any(c.rowspan > 1 or c.colspan > 1 for row in t.rows for c in row.cells)
           or max((len(row.cells) for row in t.rows), default=0) > 6 for t in page_tables):
        reasons.append('complex_table_headers_or_columns')
    if any(len(t.rows) >= 3 and all(
            sum(bool(c.text) for c in row.cells) >= 2 and
            not any(re.fullmatch(r'[+-]?\d[\d\s.,]*(?:\s*[₸₽$€]|\s*(?:KZT|USD|EUR|RUB))?', c.text)
                    for c in row.cells if c.text)
            for row in t.rows[:2]) for t in page_tables):
        reasons.append('multirow_table_header')
    if len([width for width in row_widths if width >= 6]) >= 3:
        reasons.append('complex_unruled_table')
    if not tables and sum(width >= 3 for width in row_widths) >= 3 and sum(
            bool(re.search(r'\d', c.text)) for c in source.cells if c.page == page_no) >= 4:
        reasons.append('ambiguous_unruled_numeric_table')
    # A page can contain a scan alongside a short digital title/footer.
    if page.get_images() and sum(ch.isalnum() for ch in text) < 80:
        reasons.append('image_with_sparse_native_text')
    return reasons


def _vision_render_pages(source, pdf, numbers):
    from .vision import RenderedPage
    dpi = min(300, max(72, int(os.getenv('VISION_RENDER_DPI', '144'))))
    max_side = min(4096, max(512, int(os.getenv('VISION_MAX_SIDE', '2400'))))
    for number in numbers:
        source.check()
        page = pdf[number - 1]
        scale = min(dpi / 72, max_side / max(page.rect.width, page.rect.height))
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False,
                             colorspace=fitz.csRGB)
        yield RenderedPage(page=number, image=pix.tobytes('png'), width=pix.width,
                           height=pix.height, page_width=page.rect.width,
                           page_height=page.rect.height)


def read_pdf(source, content):
    if not content.startswith(b'%PDF'):
        raise DocumentError('Содержимое файла не соответствует PDF')
    with fitz.open(stream=content, filetype='pdf') as pdf:
        if pdf.needs_pass or len(pdf) > MAX_PAGES:
            raise DocumentError(f'PDF защищён паролем или превышает {MAX_PAGES} страниц')
        source.pages = len(pdf)
        native_started = time.perf_counter()
        required = []
        for number, page in enumerate(pdf, 1):
            source.check()
            reasons = _pdf_native_page(source, page, number)
            if not page.get_images() and sum(ch.isalnum() for ch in page.get_text('text')) >= 20:
                source.routing.setdefault('nativeCrossCheckPages', []).append(number)
            if reasons:
                required.append(number)
                source.routing['reasons'].extend(f'page-{number}:{reason}' for reason in reasons)
        source.routing.update(mandatory=bool(required), required_pages=required,
                              timingsMs={'native': round((time.perf_counter() - native_started) * 1000),
                                         'render': 0, 'vision': 0, 'fusion': 0})
        force = os.getenv('VISION_PDF_MODE', 'always').casefold() == 'always'
        enabled = os.getenv('VISION_ENABLED', 'true').casefold() not in {'false', '0', 'no'}
        requested = list(range(1, len(pdf) + 1)) if force else required
        if force:
            source.routing.update(mandatory=True, required_pages=requested)
            source.routing['reasons'].append('vision_pdf_mode_always')
        if not requested:
            return
        if not enabled:
            source.routing.update(available=False, status='fallback')
            issue = 'Чтение PDF через сервис отключено; текстовый слой сохранён для проверки'
            source.routing['issues'].append(issue)
            source.warnings.append(issue)
        else:
            from .fusion import DocumentFusion
            from .vision import create_vision_service
            limit = min(MAX_PAGES, max(1, int(os.getenv('VISION_MAX_PAGES', '50'))))
            selected = requested[:limit]
            source.routing['requested'] = True
            try:
                render_started = time.perf_counter()
                rendered = list(_vision_render_pages(source, pdf, selected))
                source.routing['timingsMs']['render'] = round((time.perf_counter() - render_started) * 1000)
                vision_started = time.perf_counter()
                result = create_vision_service().analyze_pages(rendered, remaining=DEADLINE - (time.monotonic() - source.started))
                source.routing['timingsMs']['vision'] = round((time.perf_counter() - vision_started) * 1000)
                source.routing['available'] = result.available
                source.warnings.extend(result.warnings)
                if result.pages:
                    fusion_started = time.perf_counter()
                    DocumentFusion().merge(source, result)
                    source.routing['timingsMs']['fusion'] = round((time.perf_counter() - fusion_started) * 1000)
                if not result.available:
                    source.routing['status'] = 'fallback'
                    # Transport errors are sanitized by the service boundary;
                    # keep their class visible without exposing source content.
                    source.routing['error'] = result.error
                    issue = 'сервис не завершил чтение PDF; доступный текст сохранён, документ требует проверки'
                    source.routing['issues'].append(issue)
                    source.warnings.append(issue)
            except (RuntimeError, OSError, ValueError) as error:
                source.routing.update(available=False, status='fallback')
                issue = f'Визуальная обработка недоступна ({type(error).__name__}); PDF требует проверки'
                source.routing['issues'].append(issue)
                source.warnings.append(issue)
            missing = sorted(set(requested) - set(source.routing['processed_pages']))
            if missing:
                source.routing['status'] = 'limited' if source.routing['used'] else 'fallback'
                issue = f'Визуальная структура не проверена для страниц: {", ".join(map(str, missing))}'
                source.routing['issues'].append(issue)
                source.warnings.append(issue)


def read_images(source, images):
    """Read image frames through OpenAI; rendering does not perform local OCR."""
    from .vision import RenderedPage, create_vision_service
    from .fusion import DocumentFusion
    from PIL import Image, ImageOps, UnidentifiedImageError

    rendered = []
    first_page = max((p.page for p in source.page_objects), default=0) + 1
    max_side = min(4096, max(512, int(os.getenv('VISION_MAX_SIDE', '2400'))))
    limit = min(MAX_PAGES, max(1, int(os.getenv('VISION_MAX_PAGES', '50'))))
    started = time.perf_counter()
    try:
        for label, content in images:
            with Image.open(BytesIO(content)) as image:
                for frame in range(getattr(image, 'n_frames', 1)):
                    source.check()
                    if len(rendered) >= limit:
                        raise DocumentError(f'Изображения превышают лимит {limit} страниц')
                    image.seek(frame)
                    if image.width * image.height > 40_000_000:
                        raise DocumentError('Изображение превышает лимит 40 миллионов пикселей')
                    picture = ImageOps.exif_transpose(image).convert('RGB')
                    width, height = picture.size
                    picture.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
                    buffer = BytesIO()
                    picture.save(buffer, format='PNG')
                    number = first_page + len(rendered)
                    source.ensure_page(number, width, height)
                    rendered.append(RenderedPage(number, buffer.getvalue(), picture.width,
                                                 picture.height, width, height))
                    source.routing.setdefault('imageSources', []).append({'page': number, 'name': label})
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as error:
        raise DocumentError('Изображение повреждено или имеет неподдерживаемый размер') from error
    if not rendered:
        return
    source.pages = len(source.page_objects)
    timings = source.routing.setdefault('timingsMs', {})
    timings['render'] = timings.get('render', 0) + round((time.perf_counter() - started) * 1000)
    requested = [p.page for p in rendered]
    source.routing.update(mandatory=True, requested=True, required_pages=requested)
    started = time.perf_counter()
    result = create_vision_service().analyze_pages(rendered, remaining=DEADLINE - (time.monotonic() - source.started))
    timings['vision'] = round((time.perf_counter() - started) * 1000)
    source.routing['available'] = result.available
    source.warnings.extend(result.warnings)
    if result.pages:
        started = time.perf_counter()
        DocumentFusion().merge(source, result)
        timings['fusion'] = round((time.perf_counter() - started) * 1000)
    missing = sorted(set(requested) - set(source.routing['processed_pages']))
    if not result.available or missing:
        source.routing.update(status='limited' if source.routing['used'] else 'fallback', error=result.error)
        issue = 'сервис не завершил чтение изображений; часть данных не удалось прочитать'
        source.routing['issues'].append(issue)
        source.warnings.append(issue)


def read_embedded_images(source, content, suffix):
    """Office media are also read by the API, never silently ignored."""
    prefix = 'word/media/' if suffix == '.docx' else 'xl/media/'
    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = sorted(n for n in archive.namelist() if n.startswith(prefix) and not n.endswith('/'))
        supported, unsupported = [], []
        seen = set()
        import hashlib
        for name in names:
            if Path(name).suffix.lower() not in IMAGE_FORMATS:
                unsupported.append(name)
                continue
            image = archive.read(name)
            digest = hashlib.sha256(image).digest()
            if digest not in seen:
                supported.append((name, image))
                seen.add(digest)
        if supported:
            read_images(source, supported)
        if unsupported:
            issue = 'Не удалось прочитать встроенные изображения неподдерживаемого формата: ' + ', '.join(unsupported)
            source.routing.update(mandatory=True, available=False, status='limited')
            source.routing['issues'].append(issue)
            source.warnings.append(issue)

def _docx_table(source, table, block, reading_order):
    locations, elements = {}, []
    for row_no, row in enumerate(table.rows, 1):
        if row_no > MAX_ROWS or len(row.cells) > 100:
            raise DocumentError(f'Таблица превышает {MAX_ROWS:,} строк / 100 колонок'.replace(',', ' '))
        for column, cell in enumerate(row.cells, 1):
            # python-docx returns the same XML cell at every occupied grid
            # location for horizontal and vertical merges.
            element = cell._tc
            locations.setdefault(element, []).append((row_no, column))
            elements.append((row_no, column, cell))
    anchors = {}
    for row_no, column, word_cell in elements:
        positions = locations[word_cell._tc]
        anchor = anchors.get(word_cell._tc)
        item = source.add('' if anchor else word_cell.text, block, row_no, column,
                          kind='table', source_method='docx-native', reading_order=reading_order,
                          merged_into=anchor.id if anchor else None,
                          rowspan=1 if anchor else max(y for y, _ in positions) - row_no + 1,
                          colspan=1 if anchor else max(x for _, x in positions) - column + 1)
        if not anchor:
            anchors[word_cell._tc] = item
    canonical = source._table_index.get(block)
    if canonical:
        canonical.header_rows = [index for index, row in enumerate(table.rows, 1)
                                 if row._tr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}trPr/'
                                                 '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader') is not None]


def read_docx(source, content):
    check_zip(content)
    doc = Document(BytesIO(content))
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for index, element in enumerate(doc.element.body):
        if element.tag.endswith('}p'):
            paragraph = Paragraph(element, doc)
            add_lines(source, paragraph.text, f'paragraph-{index}',
                      source_method='docx-native', reading_order=index)
            block = source._block_index.get(f'paragraph-{index}')
            if block and paragraph.style and paragraph.style.name.startswith(('Title', 'Heading')):
                block.type = 'title'
        elif element.tag.endswith('}tbl'):
            _docx_table(source, Table(element, doc), f'table-{index}', index)
    visited = set()
    for index, section in enumerate(doc.sections):
        for kind in ('header', 'footer'):
            for region in (getattr(section, kind), getattr(section, f'first_page_{kind}'),
                           getattr(section, f'even_page_{kind}')):
                part = str(region.part.partname)
                if part in visited:
                    continue
                visited.add(part)
                prefix = f'{kind}-{index}-{len(visited)}'
                for element_index, element in enumerate(region._element):
                    order = -1000 + element_index if kind == 'header' else 100000 + element_index
                    if element.tag.endswith('}p'):
                        add_lines(source, Paragraph(element, region).text, f'{prefix}-paragraph-{element_index}',
                                  source_method='docx-native', reading_order=order)
                    elif element.tag.endswith('}tbl'):
                        _docx_table(source, Table(element, region), f'{prefix}-table-{element_index}', order)
    if any(c.rowspan > 1 or c.colspan > 1 for c in source.cells):
        source.routing['reasons'].append('native_docx_merged_cells')
    source.routing['reasons'].append('docx_native_structure')

def _xlsx_merges(book, sheet):
    """Read merge declarations without loading a second full workbook into RAM."""
    from openpyxl.utils.cell import range_boundaries
    merges = []
    with book._archive.open(sheet._worksheet_path) as stream:
        for _, element in ElementTree.iterparse(stream, events=('end',)):
            if element.tag.endswith('}mergeCell'):
                bounds = range_boundaries(element.attrib['ref'])
                if bounds[2] > 100 or bounds[3] > MAX_ROWS:
                    raise DocumentError(f'Лист превышает {MAX_ROWS:,} строк / 100 колонок'.replace(',', ' '))
                merges.append(bounds)
            element.clear()
    return merges


def read_excel(source, content, suffix):
    if suffix == '.xlsx':
        import openpyxl
        check_zip(content); book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            if len(book.worksheets) > MAX_SHEETS: raise DocumentError(f'Превышен лимит {MAX_SHEETS} листов')
            for sheet in book:
                from openpyxl.utils.cell import get_column_letter
                merges = _xlsx_merges(book, sheet)
                if sum((right - left + 1) * (bottom - top + 1)
                       for left, top, right, bottom in merges) > MAX_CELLS:
                    raise DocumentError(f'Объединённые ячейки превышают лимит {MAX_CELLS} ячеек')
                merge_positions = {(row, column): (left, top, right, bottom)
                                   for left, top, right, bottom in merges
                                   for row in range(top, bottom + 1)
                                   for column in range(left, right + 1)}
                merge_widths = {}
                for left, top, right, bottom in merges:
                    for row in range(top, bottom + 1):
                        merge_widths[row] = max(merge_widths.get(row, 0), right)
                anchors = {}
                sheet.reset_dimensions()
                def add_row(row_no, values):
                    if row_no > MAX_ROWS or len(values) > 100: raise DocumentError(f'Лист превышает {MAX_ROWS:,} строк / 100 колонок'.replace(',', ' '))
                    values = tuple(values) + (None,) * max(0, merge_widths.get(row_no, 0) - len(values))
                    for col, value in enumerate(values, 1):
                        merge = merge_positions.get((row_no, col))
                        anchor = anchors.get((merge[1], merge[0])) if merge else None
                        item = source.add(value, f'sheet-{sheet.title}', row_no, col,
                                          sheet=sheet.title, kind='table', source_method='xlsx-native',
                                          coordinate=f'{get_column_letter(col)}{row_no}',
                                          merged_into=anchor.id if anchor else None,
                                          rowspan=merge[3] - merge[1] + 1 if merge and not anchor else 1,
                                          colspan=merge[2] - merge[0] + 1 if merge and not anchor else 1)
                        if merge and not anchor:
                            anchors[(merge[1], merge[0])] = item
                row_no = 0
                for row_no, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    add_row(row_no, values)
                for trailing in range(row_no + 1, max(merge_widths, default=row_no) + 1):
                    add_row(trailing, ())
            source.warnings.append('Excel: формулы читаются из сохранённого результата; пустой кэш остаётся пустым')
        finally: book.close()
    else:
        import xlrd
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            if book.nsheets > MAX_SHEETS: raise DocumentError(f'Превышен лимит {MAX_SHEETS} листов')
            for sheet in book.sheets():
                if sheet.nrows > MAX_ROWS or sheet.ncols > 100: raise DocumentError(f'Лист превышает {MAX_ROWS:,} строк / 100 колонок'.replace(',', ' '))
                source.tables += 1
                for row in range(sheet.nrows):
                    for col in range(sheet.ncols):
                        item = sheet.cell(row, col); value = xlrd.xldate.xldate_as_datetime(item.value, book.datemode).date().isoformat() if item.ctype == 3 else item.value
                        source.add(value, f'sheet-{sheet.name}', row + 1, col + 1,
                                   sheet=sheet.name, kind='table', source_method='xls-native')
        finally: book.release_resources()

def read_text(source, content, suffix):
    text = decode(content)
    if len(text) > MAX_CHARS: raise DocumentError(f'Превышен лимит {MAX_CHARS} символов')
    if suffix == '.json':
        value = json.loads(text); value = value.get('proposal', value) if isinstance(value, dict) else value
        if isinstance(value, list): value = {'items': value}
        if not isinstance(value, dict): raise DocumentError('JSON должен содержать объект предложения или массив позиций')
        for row, (key, item) in enumerate(value.items(), 1):
            if isinstance(item, list) and all(isinstance(entry, dict) for entry in item):
                keys = list(dict.fromkeys(k for entry in item for k in entry)); source.tables += 1
                for col, name in enumerate(keys, 1): source.add(name, f'json-{key}', 1, col, kind='table')
                for idx, entry in enumerate(item, 2):
                    for col, name in enumerate(keys, 1): source.add(entry.get(name), f'json-{key}', idx, col, kind='table')
            elif not isinstance(item, (dict, list)): source.add(f'{key}: {item if item is not None else ""}', 'json-fields', row)
            else: source.warnings.append(f'JSON: неподдерживаемая вложенная структура {key}')
    elif suffix in {'.csv', '.tsv'}:
        try: dialect = csv.Sniffer().sniff(text[:8192], delimiters=';\t,|')
        except csv.Error: dialect = csv.excel_tab if suffix == '.tsv' else csv.excel
        source.tables += 1
        for row, values in enumerate(csv.reader(StringIO(text), dialect), 1):
            for col, value in enumerate(values, 1): source.add(value, 'text-table', row, col, kind='table')
    else: add_lines(source, text, 'text', split_blocks=True)

def read_document(content: bytes, filename: str) -> Source:
    suffix = Path(filename).suffix.lower()
    if suffix not in FORMATS: raise DocumentError('Неподдерживаемый формат. Используйте PDF, изображения, DOCX, XLSX/XLS, CSV/TSV, TXT или JSON.')
    if not content or len(content) > MAX_BYTES: raise DocumentError(f'Файл пуст или превышает {MAX_BYTES // 1024**2} МБ')
    source = Source(filename)
    source.routing['format'] = suffix.lstrip('.')
    if suffix == '.pdf': read_pdf(source, content)
    elif suffix in IMAGE_FORMATS: read_images(source, [(filename, content)])
    elif suffix == '.docx': read_docx(source, content)
    elif suffix in {'.xlsx', '.xls'}: read_excel(source, content, suffix)
    else: read_text(source, content, suffix)
    if suffix in {'.docx', '.xlsx'}:
        read_embedded_images(source, content, suffix)
    if suffix == '.xls':
        # xlrd reads cells but cannot expose embedded pictures. Never certify a
        # complete visual review of an old binary workbook that we cannot inspect.
        source.warnings.append('XLS: встроенные изображения не проверены. Для их чтения загрузите XLSX или PDF.')
        source.routing['issues'].append('xls_embedded_images_unverified')
        source.routing['embeddedImagesReviewRequired'] = True
    block_order = {block.id: index for index, block in enumerate(source.blocks)}
    source.cells.sort(key=lambda cell: (cell.page or 0, cell.reading_order,
                                       block_order[cell.block], cell.row, cell.cell))
    source.reindex()
    return source
