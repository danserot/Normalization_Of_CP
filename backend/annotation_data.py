"""Human review data and inference-compatible training examples. No model calls."""
import hashlib
import json
import os
import re
from dataclasses import fields as dataclass_fields
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Literal

import fitz
from docx import Document
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from pydantic import Field

from .extraction import Cell, DocumentError, Source, check_zip, decode, read_document
from .rules import FIELDS, HEADERS, extract_rules, normalized
from .universal import LAYOUT_SYSTEM, REQUISITES_SYSTEM, Layout, Strict, catalog, metadata_cells, numeric_quote, requisite_batches


class MarkedField(Strict):
    field: str = Field(max_length=100)
    state: str = Field(pattern='^(pending|found|missing)$')
    cell: str = Field(default='', max_length=100)
    value: str = Field(default='', max_length=2000)


class MarkedTable(Layout):
    block: str = Field(max_length=200)
    reviewed: bool = False
    componentMode: Literal['additive', 'alternative'] = 'additive'
    unitPriceColumn: int = Field(default=0, ge=0, le=100)
    lineTotalColumn: int = Field(default=0, ge=0, le=100)


class Annotation(Strict):
    fields: list[MarkedField] = Field(max_length=30)
    tables: list[MarkedTable] = Field(max_length=100)
    issues: list[str] = Field(default_factory=list, max_length=30)
    group: str = Field(default='', max_length=200)
    notes: str = Field(default='', max_length=5000)


class SaveAnnotation(Strict):
    revision: int = Field(ge=1)
    status: str = Field(pattern='^(draft|reviewed)$')
    annotation: Annotation


def source_from_payload(payload, filename):
    """Re-open stored sources, including snapshots made before canonical fusion."""
    supported = {field.name for field in dataclass_fields(Cell)}
    aliases = {'sourceMethod': 'source_method', 'visionAgreement': 'vision_agreement',
               'readingOrder': 'reading_order', 'rowSpan': 'rowspan', 'columnSpan': 'colspan'}
    cells = []
    for raw in payload['cells']:
        cell = {aliases.get(key, key): value for key, value in raw.items()}
        if 'cell' not in cell and 'column' in cell:
            cell['cell'] = cell['column']
        cells.append(Cell(**{key: value for key, value in cell.items() if key in supported}))
    source = Source(file=filename, cells=cells)
    source.warnings = list(payload.get('warnings', []))
    source.pages = max((cell.page or 0 for cell in cells), default=0)
    source.tables = len({cell.block for cell in cells if cell.kind == 'table'})
    source.chars = sum(len(cell.text) for cell in cells)
    if hasattr(source, 'routing'):
        source.routing = dict(payload.get('routing', {}))
    return source


def draft_annotation(source, proof):
    fields = []
    for key in FIELDS:
        candidate = proof.get(key)
        cell = next((c for c in source.cells if isinstance(candidate, dict)
                     and c.block == candidate.get('block') and c.row == candidate.get('row')
                     and c.cell == candidate.get('cell')), None)
        value = str(candidate.get('value', '')) if isinstance(candidate, dict) else ''
        if key == 'documentTotal' and cell and isinstance(candidate, dict):
            # Keep the original numeric spelling, including grouping/decimals.
            value = next((match.group().strip() for match in re.finditer(r'[+]?\d[\d\s\u00a0\u202f.,]*', cell.text)
                          if numeric_quote(match.group().strip()) == candidate.get('value')), '')
        if not cell or not value or value.casefold() not in cell.text.casefold():
            cell, value = None, ''
        fields.append({'field': key, 'state': 'pending', 'cell': cell.id if cell else '', 'value': value})
    tables = []
    for block in catalog(source):
        cells = [c for c in source.cells if c.block == block['block']]
        width = max(c.cell for c in cells)
        if width < 2:
            continue
        # Suggestions only; the reviewer must explicitly classify every block.
        header = next((row for row in block['rows'] if any(normalized(c[2]) in HEADERS['name'] for c in row['cells'])), block['rows'][0])
        roles = {role: next((c for c in header['cells'] if normalized(c[2]) in names), None) for role, names in HEADERS.items()}
        components = []
        price, total = roles['unitPrice'], roles['lineTotal']
        if price or total:
            label_cell = price or total
            components = [{'label': label_cell[2], 'labelCell': label_cell[0],
                           'priceColumn': price[1] if price else 0, 'totalColumn': total[1] if total else 0}]
        tables.append({'block': block['block'], 'reviewed': False, 'isItems': bool(roles['name']),
                       'firstRow': min(header['row'] + 1, block['lastRow']), 'lastRow': block['lastRow'],
                       'nameColumn': roles['name'][1] if roles['name'] else 0,
                       'quantityColumn': roles['quantity'][1] if roles['quantity'] else 0,
                       'unitColumn': roles['unit'][1] if roles['unit'] else 0,
                       'components': components, 'extras': [], 'componentMode': 'additive',
                       'unitPriceColumn': 0, 'lineTotalColumn': 0})
    return {'fields': fields, 'tables': tables, 'issues': [], 'group': '', 'notes': ''}


def word_preview(content):
    """Safe text/table preview. Embedded raster images are intentionally ignored."""
    check_zip(content)
    doc = Document(BytesIO(content))

    def paragraph_html(paragraph):
        parts = []
        for element in paragraph._p:
            if element.tag.endswith('}hyperlink'):
                parts.append(escape(''.join(node.text or '' for node in element.xpath('.//w:t'))))
                continue
            if not element.tag.endswith('}r'):
                continue
            run = Run(element, paragraph)
            text = escape(run.text).replace('\n', '<br>')
            if run.bold:
                text = '<strong>' + text + '</strong>'
            if run.italic:
                text = '<em>' + text + '</em>'
            parts.append(text)
        if not parts:
            parts.append(escape(paragraph.text))
        return '<p>' + ''.join(parts) + '</p>'

    def table_html(table):
        def cell_html(cell):
            return ''.join(paragraph_html(Paragraph(element, cell)) if element.tag.endswith('}p')
                           else table_html(WordTable(element, cell)) if element.tag.endswith('}tbl') else ''
                           for element in cell._tc)
        return '<table>' + ''.join('<tr>' + ''.join('<td>' + cell_html(cell) + '</td>' for cell in row.cells)
                                  + '</tr>' for row in table.rows) + '</table>'

    parts = []
    for section in doc.sections:
        parts.extend(paragraph_html(p) for p in section.header.paragraphs if p.text)
    for element in doc.element.body:
        if element.tag.endswith('}p'):
            parts.append(paragraph_html(Paragraph(element, doc)))
        elif element.tag.endswith('}tbl'):
            parts.append(table_html(WordTable(element, doc)))
    for section in doc.sections:
        parts.extend(paragraph_html(p) for p in section.footer.paragraphs if p.text)
    return '<!doctype html><html lang="ru"><meta charset="utf-8"><style>body{font:15px/1.5 Arial;margin:24px;color:#243442;background:white}table{border-collapse:collapse;width:100%;margin:20px 0}td{border:1px solid #ccd5df;padding:8px;vertical-align:top}p{white-space:pre-wrap}img{max-width:100%;height:auto}</style><body>' + ''.join(parts) + '</body></html>'


def prepare_annotation(content, filename):
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        with fitz.open(stream=content, filetype='pdf') as pdf:
            if pdf.needs_pass or not 0 < len(pdf) <= 50:
                raise DocumentError('PDF защищён паролем или превышает 50 страниц')
            preview = {'kind': 'pdf', 'pages': len(pdf)}
    elif suffix == '.docx':
        preview = {'kind': 'word', 'html': word_preview(content)}
    elif suffix in {'.txt', '.json'}:
        text = decode(content)
        if len(text) > 200_000:
            raise DocumentError('Превышен лимит 200 000 символов')
        preview = {'kind': 'text', 'text': text}
    else:
        preview = {'kind': 'table'}
    try:
        source = read_document(content, filename)
    except DocumentError as error:
        if preview['kind'] not in {'pdf', 'word', 'text'}:
            raise
        return {'cells': [], 'warnings': [str(error)], 'parseError': str(error), 'preview': preview,
                'annotation': draft_annotation(Source(filename), {})}
    proposal, proof, warnings, _ = extract_rules(source)
    return {'cells': source.public(), 'warnings': list(dict.fromkeys([*source.warnings, *warnings])),
            'routing': getattr(source, 'routing', {}),
            'canonicalDocument': source.structure() if hasattr(source, 'structure') else {},
            'parseError': '', 'preview': preview,
            'suggestion': proposal, 'annotation': draft_annotation(source, proof)}


def validate_annotation(annotation, payload, filename, reviewed):
    source = source_from_payload(payload, filename)
    by_id = {c.id: c for c in source.cells}
    if len({f.field for f in annotation.fields}) != len(annotation.fields) or {f.field for f in annotation.fields} != set(FIELDS):
        raise ValueError('Разметка должна содержать каждое поле ровно один раз')
    blocks = {b['block']: b for b in catalog(source) if max(c.cell for c in source.cells if c.block == b['block']) >= 2}
    if len({t.block for t in annotation.tables}) != len(annotation.tables) or {t.block for t in annotation.tables} != set(blocks):
        raise ValueError('Проверьте все табличные блоки документа')
    if reviewed and (payload.get('parseError') or not source.cells):
        raise ValueError('Сначала исправьте ошибку чтения документа; без ячеек обучение невозможно')
    if reviewed and not annotation.group.strip():
        raise ValueError('Укажите группу документов / шаблон поставщика')
    if not reviewed:
        return
    for field in annotation.fields:
        if reviewed and field.state == 'pending':
            raise ValueError(f'Проверьте поле {field.field}: значение или «Не указано»')
        if field.state == 'found':
            cell = by_id.get(field.cell)
            if not cell or not field.value.strip() or field.value not in cell.text:
                raise ValueError(f'{field.field}: значение должно быть точной цитатой выбранной ячейки')
            if field.field == 'documentTotal' and numeric_quote(field.value) is None:
                raise ValueError('Общий итог: выделите только число из исходной ячейки, без пояснений')
        if field.state == 'missing' and (field.cell or field.value):
            raise ValueError(f'{field.field}: у отсутствующего поля не должно быть значения')
    for table in annotation.tables:
        if reviewed and not table.reviewed:
            raise ValueError(f'Подтвердите тип блока {table.block}')
        if not table.isItems:
            continue
        block = blocks[table.block]
        width = max(c.cell for c in source.cells if c.block == table.block)
        if not block['firstRow'] <= table.firstRow <= table.lastRow <= block['lastRow']:
            raise ValueError(f'{table.block}: неверный диапазон строк')
        if not 1 <= table.nameColumn <= width:
            raise ValueError(f'{table.block}: выберите колонку наименования')
        columns = [table.nameColumn, table.quantityColumn, table.unitColumn,
                   table.unitPriceColumn, table.lineTotalColumn]
        for component in table.components:
            columns.extend([component.priceColumn, component.totalColumn])
            if not component.priceColumn and not component.totalColumn:
                raise ValueError('Для стоимости нужна колонка цены или суммы')
        columns.extend(extra.column for extra in table.extras)
        nonzero = [c for c in columns if c]
        if any(c > width for c in columns) or len(nonzero) != len(set(nonzero)):
            raise ValueError(f'{table.block}: колонки не должны повторяться или выходить за границы')
        for labelled in [*table.components, *table.extras]:
            cell = by_id.get(labelled.labelCell)
            if not cell or not labelled.label.strip() or labelled.label not in cell.text:
                raise ValueError('Название стоимости / дополнительной колонки должно ссылаться на заголовок в источнике')
    # Requisites are requested outside item blocks by the production pipeline.
    metadata_ids = {c.id for c in metadata_cells(source, [t for t in annotation.tables if t.isItems])}
    if any(f.state == 'found' and f.cell not in metadata_ids for f in annotation.fields):
        raise ValueError('Реквизит попал в диапазон товарных строк. Исправьте начало/конец таблицы: исключите заголовки и итоги')
    previous_ids = set()
    for block in catalog(source):
        table = next((t for t in annotation.tables if t.block == block['block']), None)
        if table is None or not table.isItems:
            continue
        visible = {cell[0] for row in block['rows'] for cell in row['cells']} | previous_ids
        if any(label.labelCell not in visible for label in [*table.components, *table.extras]):
            raise ValueError('Выбранный заголовок отсутствует в контексте этого блока и предыдущей таблицы. Проверьте блок и ссылку на заголовок')
        previous_ids = {label.labelCell for label in [*table.components, *table.extras]}


def training_examples(annotation, payload, filename):
    source = source_from_payload(payload, filename)
    layouts = {t.block: t for t in annotation.tables}
    result, previous = [], None

    def add(system, user, answer):
        result.append({'system': system, 'instruction': json.dumps(user, ensure_ascii=False, separators=(',', ':')),
                       'input': '', 'output': json.dumps(answer, ensure_ascii=False, separators=(',', ':'))})

    for block in catalog(source):
        table = layouts.get(block['block'])
        if table is None:
            continue
        answer = table.model_dump(include=set(Layout.model_fields))
        if not table.isItems:
            answer.update(nameColumn=0, quantityColumn=0, unitColumn=0, components=[], extras=[],
                          firstRow=block['firstRow'], lastRow=block['lastRow'])
        add(LAYOUT_SYSTEM, {**block, 'previousLayout': previous}, Layout.model_validate(answer).model_dump())
        if table.isItems:
            previous = answer
    for batch in requisite_batches(metadata_cells(source, [t for t in annotation.tables if t.isItems])):
        ids = {c[0] for c in batch}
        fields = [{'field': f.field, 'cell': f.cell, 'value': f.value}
                  for f in annotation.fields if f.state == 'found' and f.cell in ids]
        add(REQUISITES_SYSTEM, batch, {'fields': fields})
    return result


def plan_training_examples(annotation, payload, filename):
    """One compact semantic plan; Python retains responsibility for all item rows."""
    from .universal import PLAN_SYSTEM, Plan, Quote, Table, plan_payload

    source = source_from_payload(payload, filename)
    answer = Plan(
        fields=[Quote(field=field.field, cell=field.cell, value=field.value)
                for field in annotation.fields if field.state == 'found'],
        tables=[Table.model_validate(table.model_dump(exclude={'isItems', 'reviewed'}))
                for table in annotation.tables if table.isItems],
        # Reviewer notes/issues are audit information, not sourced model answers.
        issues=[],
    )
    model_input = plan_payload(source)
    visible_ids = set()

    def collect_ids(value):
        if isinstance(value, dict):
            for key, nested in value.items():
                if key == 'cells':
                    visible_ids.update(cell[0] for cell in nested)
                else:
                    collect_ids(nested)
        elif isinstance(value, list):
            for nested in value:
                collect_ids(nested)

    collect_ids(model_input)
    references = {field.cell for field in answer.fields}
    references.update(column.labelCell for table in answer.tables
                      for column in [*table.components, *table.extras])
    if not references <= visible_ids:
        raise DocumentError('Проверенные ссылки отсутствуют в компактном контексте модели; уточните структуру документа')
    instruction = json.dumps(model_input, ensure_ascii=False, separators=(',', ':'))
    context_limit = int(os.getenv('SEMANTIC_CONTEXT_CHARS', os.getenv('SEMANTIC_CONTEXT_CHAR_LIMIT', '60000')))
    if len(instruction) > context_limit:
        raise DocumentError('Документ превышает лимит semantic context; экспорт с обрезанными ссылками запрещён')
    return [{'system': PLAN_SYSTEM,
             'instruction': instruction,
             'input': '', 'output': answer.model_dump_json()}]


def group_split(group):
    bucket = int(hashlib.sha256(group.strip().casefold().encode()).hexdigest()[:8], 16) % 100
    return 'train' if bucket < 80 else 'val' if bucket < 90 else 'test'
