"""Contract tests for visual layout, native authority and adaptive routing."""
from io import BytesIO
from unittest.mock import Mock

import fitz
import pytest
from docx import Document
from openpyxl import Workbook

from backend.canonical import CanonicalDocument, CanonicalCell, NativeWord
from backend.extraction import DocumentError, read_document
from backend.fusion import DocumentFusion, parse_table_html
from backend.vision import VisionBlock, VisionPage, VisionResult


def _pdf(text=True, table=False):
    with fitz.open() as pdf:
        page = pdf.new_page(width=300, height=400)
        if text:
            page.insert_text((30, 30), 'Proposal document with accurate native text', fontsize=9)
        if table:
            for x in (30, 160, 270):
                page.draw_line((x, 80), (x, 140))
            for y in (80, 110, 140):
                page.draw_line((30, y), (270, y))
            for x, y, value in ((40, 100, 'Name'), (170, 100, 'Price'),
                                (40, 130, 'Cable'), (170, 130, '1 250 000')):
                page.insert_text((x, y), value, fontsize=9)
        return pdf.tobytes()


def _source_with_native_grid():
    source = CanonicalDocument('price.pdf')
    source.ensure_page(1, 200, 100)
    values = [('Name', [0, 0, 100, 30]), ('Price', [100, 0, 200, 30]),
              ('Cable', [0, 30, 100, 60]), ('1 250 000', [100, 30, 200, 60])]
    for index, (text, box) in enumerate(values):
        source.add(text, 'native-table', index // 2 + 1, index % 2 + 1,
                   page=1, kind='table', bbox=box, source_method='pdf-native')
        source.native_words.append(NativeWord(f'w{index}', text, 1, box, index // 2, 0, index))
    return source


def test_canonical_hierarchy_references_flat_cells_and_keeps_training_constructor():
    source = _source_with_native_grid()
    rebuilt = CanonicalDocument('price.pdf', cells=[CanonicalCell(**c) for c in source.public()])
    assert rebuilt.table_objects[0].id == rebuilt.cells[0].block
    assert rebuilt.table_objects[0].rows[1].cells[1] is rebuilt.cells[3]
    assert rebuilt.page_objects[0].block_ids == ['native-table']
    assert source.structure()['tables'][0]['rows'][1]['cell_ids'] == ['c2', 'c3']
    assert source.cells[3].column == 2


def test_html_spans_preserve_logical_grid_and_do_not_repeat_merged_text():
    cells = parse_table_html('<table><tr><th rowspan="2">Name</th><th colspan="2">Cost</th></tr>'
                             '<tr><th>Equipment</th><th>Work</th></tr>'
                             '<tr><td>Cable</td><td>100</td><td>20</td></tr></table>')
    assert [(c.text, c.row, c.column) for c in cells[:4]] == [
        ('Name', 1, 1), ('Cost', 1, 2), ('Equipment', 2, 2), ('Work', 2, 3)]
    assert cells[0].rowspan == 2 and cells[1].colspan == 2


@pytest.mark.parametrize('cell_boxes', [True, False])
@pytest.mark.parametrize('native_row_offset', [0, 12])
def test_native_number_wins_over_visual_typo_with_explicit_conflict(cell_boxes, native_row_offset):
    source = _source_with_native_grid()
    for cell in source.cells:
        cell.row += native_row_offset
    source.reindex()
    block = VisionBlock('visual-table', 'table', [0, 0, 200, 60], html=(
        '<table><tr><th>Name</th><th>Price</th></tr>'
        '<tr><td>Cable</td><td>1 250 800</td></tr></table>'),
        cell_boxes=[c.bbox for c in source.cells] if cell_boxes else [])
    DocumentFusion().merge(source, VisionResult([VisionPage(1, [block])]))
    price = next(c for c in source.cells if c.block == 'visual-table' and c.row == 2 and c.cell == 2)
    assert price.text == '1 250 000'
    assert price.source_method == 'pdf-native+vision'
    assert price.vision_agreement is False
    assert any('расходится' in warning for warning in source.warnings)
    assert source.tables == 1
    assert len(source.cells) == 4
    assert len({c.id for c in source.cells}) == len(source.cells)


def test_vision_only_table_retains_spans_and_empty_placeholders():
    source = CanonicalDocument('scan.pdf')
    source.ensure_page(1, 200, 100)
    block = VisionBlock('table', 'table', [0, 0, 200, 100], html=(
        '<table><tr><th colspan="2">Costs</th></tr>'
        '<tr><td>Cable</td><td>100</td></tr></table>'))
    DocumentFusion().merge(source, VisionResult([VisionPage(1, [block])]))
    header = next(c for c in source.cells if c.text == 'Costs')
    placeholder = next(c for c in source.cells if c.row == 1 and c.cell == 2)
    assert header.colspan == 2 and placeholder.text == ''
    assert placeholder.merged_into == header.id
    assert source.table_objects[0].header_rows == [1]
    assert source.routing['used'] and source.routing['status'] == 'vision'


def test_visual_failure_cannot_delete_native_table_or_mark_page_processed():
    source = _source_with_native_grid()
    block = VisionBlock('broken', 'table', [0, 0, 200, 60], html='<table></table>')
    DocumentFusion().merge(source, VisionResult([VisionPage(1, [block])]))
    assert [c.text for c in source.cells] == ['Name', 'Price', 'Cable', '1 250 000']
    assert source.routing['processed_pages'] == []
    assert 'invalid_visual_table' in source.routing['issues']


def test_simple_native_pdf_skips_vision_and_uses_individual_table_boxes(monkeypatch):
    service = Mock()
    monkeypatch.setattr('backend.vision.create_vision_service', lambda: service)
    source = read_document(_pdf(table=True), 'simple.pdf')
    service.analyze_pages.assert_not_called()
    table = source.table_objects[0]
    assert len(table.rows) == 2
    assert all(c.bbox != table.bbox for row in table.rows for c in row.cells)
    assert source.native_words and source.page_objects[0].width == 300
    assert all(c.source_method == 'pdf-native' for c in source.cells)
    assert source.routing['status'] == 'native'


@pytest.mark.parametrize('rotation', [90, 180, 270])
def test_pdf_rotation_coordinates_match_render_and_fusion_preserves_native_values(monkeypatch, rotation):
    monkeypatch.setenv('VISION_ENABLED', 'false')
    content = _pdf(table=True)
    baseline = read_document(content, 'price.pdf')
    with fitz.open(stream=content, filetype='pdf') as pdf:
        pdf[0].set_rotation(rotation)
        content = pdf.tobytes()
    with fitz.open(stream=content, filetype='pdf') as pdf:
        page = pdf[0]
        page.set_rotation(0)
        native_box = page.find_tables().tables[0].rows[1].cells[1]
        page.set_rotation(rotation)
        expected_box = list(fitz.Rect(native_box) * page.rotation_matrix)
        expected_size = (page.rect.width, page.rect.height)
    source = read_document(content, 'price.pdf')
    price = next(c for c in source.cells if c.kind == 'table' and c.text == '1 250 000')
    assert price.bbox == pytest.approx(expected_box)
    assert source.page_objects[0].rotation == rotation
    assert (source.page_objects[0].width, source.page_objects[0].height) == expected_size
    assert [c.text for c in source.cells] == [c.text for c in baseline.cells]
    assert all(0 <= w.bbox[0] < w.bbox[2] <= expected_size[0]
               and 0 <= w.bbox[1] < w.bbox[3] <= expected_size[1] for w in source.native_words)
    block = VisionBlock('visual-price', 'text', price.bbox, text='1 250 800')
    DocumentFusion().merge(source, VisionResult([VisionPage(1, [block])]))
    fused = next(c for c in source.cells if c.block == 'visual-price')
    assert fused.text == '1 250 000' and fused.vision_agreement is False
    assert source.page_objects[0].rotation == rotation


def test_scan_routes_to_vision_and_rendering_obeys_bound(monkeypatch):
    monkeypatch.setenv('VISION_ENABLED', 'true')
    monkeypatch.setenv('VISION_MAX_SIDE', '600')
    service = Mock()
    service.analyze_pages.return_value = VisionResult([VisionPage(1, [
        VisionBlock('title', 'title', [10, 10, 290, 40], text='Commercial proposal'),
        VisionBlock('table', 'table', [10, 60, 290, 200], html=(
            '<table><tr><th>Name</th><th>Quantity</th><th>Price</th></tr>'
            '<tr><td>Cable</td><td>2</td><td>100</td></tr></table>')),
    ])])
    monkeypatch.setattr('backend.vision.create_vision_service', lambda: service)
    source = read_document(_pdf(text=False), 'scan.pdf')
    rendered = service.analyze_pages.call_args.args[0][0]
    assert max(rendered.width, rendered.height) <= 600
    assert rendered.page_width == 300 and rendered.image.startswith(b'\x89PNG')
    assert source.routing['mandatory'] and source.routing['used']
    assert source.routing['processed_pages'] == [1]
    assert any(c.text == 'Cable' and c.method == 'paddleocr-vl' for c in source.cells)


def test_scanned_pdf_failure_is_explicit(monkeypatch):
    monkeypatch.setenv('VISION_ENABLED', 'false')
    source = read_document(_pdf(text=False), 'scan.pdf')
    assert not source.text()
    assert source.routing['available'] is False
    assert source.routing['mandatory'] and source.routing['status'] == 'fallback'
    assert source.routing['issues'] and source.warnings


def test_ambiguous_borderless_table_routes_and_retains_native_on_service_failure(monkeypatch):
    with fitz.open() as pdf:
        page = pdf.new_page(width=400, height=400)
        for row, values in enumerate([['Name', 'Quantity', 'Price', 'Total'],
                                      ['Cable', '2', '100', '200'], ['Pipe', '3', '80', '240']]):
            for col, value in enumerate(values):
                page.insert_text((30 + col * 90, 30 + row * 30), value, fontsize=9)
        content = pdf.tobytes()
    monkeypatch.setenv('VISION_ENABLED', 'true')
    service = Mock()
    service.analyze_pages.return_value = VisionResult(available=False, error='offline')
    monkeypatch.setattr('backend.vision.create_vision_service', lambda: service)
    source = read_document(content, 'borderless.pdf')
    assert source.routing['mandatory'] and source.routing['status'] == 'fallback'
    assert source.routing['error'] == 'offline'
    assert 'Cable' in source.text() and 'Pipe' in source.text()
    assert any('ambiguous_unruled_numeric_table' in r for r in source.routing['reasons'])
    assert source.routing['issues']


def test_xlsx_keeps_merged_native_headers_coordinates_and_never_uses_vision(monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells('B1:C1')
    sheet['A1'], sheet['B1'] = 'Name', 'Costs'
    sheet.append(['', 'Equipment', 'Work'])
    sheet.append(['Cable', 100, 20])
    output = BytesIO()
    workbook.save(output)
    service = Mock()
    monkeypatch.setattr('backend.vision.create_vision_service', lambda: service)
    source = read_document(output.getvalue(), 'merged.xlsx')
    origin = next(c for c in source.cells if c.coordinate == 'B1')
    placeholder = next(c for c in source.cells if c.coordinate == 'C1')
    assert origin.colspan == 2 and placeholder.merged_into == origin.id
    assert origin.sheet == sheet.title and origin.source_method == 'xlsx-native'
    service.analyze_pages.assert_not_called()


def test_docx_preserves_header_footer_tables_and_merged_cells(monkeypatch):
    doc = Document()
    doc.add_paragraph('Commercial proposal')
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = 'Name'
    table.cell(0, 1).merge(table.cell(0, 2)).text = 'Costs'
    table.cell(1, 0).text = 'Cable'
    table.cell(1, 1).text = '100'
    table.cell(1, 2).text = '20'
    doc.sections[0].header.add_table(rows=1, cols=1, width=100).cell(0, 0).text = 'Supplier header'
    doc.sections[0].footer.paragraphs[0].text = 'Footer terms'
    output = BytesIO()
    doc.save(output)
    service = Mock()
    monkeypatch.setattr('backend.vision.create_vision_service', lambda: service)
    source = read_document(output.getvalue(), 'merged.docx')
    assert source.tables == 2
    assert len([c for c in source.cells if c.text == 'Costs']) == 1
    assert next(c for c in source.cells if c.text == 'Costs').colspan == 2
    assert any(c.text == 'Supplier header' and c.block.startswith('header-') for c in source.cells)
    assert any(c.text == 'Footer terms' and c.block.startswith('footer-') for c in source.cells)
    assert all(c.source_method == 'docx-native' for c in source.cells)
    service.analyze_pages.assert_not_called()
