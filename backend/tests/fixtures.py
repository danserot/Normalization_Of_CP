"""Synthetic, non-personal fixtures shared by unit tests and HTTP benchmarks."""
import json
from io import BytesIO
from pathlib import Path

import fitz
import openpyxl
import xlwt
from docx import Document
from PIL import Image, ImageDraw, ImageFont

META = ['Коммерческое предложение № КП-42 от 29.09.2026', 'Клиент: ТОО Альфа',
        'Контакт: manager@example.test', 'Действительно до: 31.10.2026', 'Поставщик: ТОО Бета',
        'Валюта: KZT', 'НДС: включён 12%', 'Скидка: 0%', 'Доставка: 0',
        'Условия оплаты: 50% предоплата', 'Срок поставки: 10 дней', 'Гарантия: 12 месяцев']
HEAD = ['Сумма', 'Ед. изм.', 'Наименование', 'Цена', 'Количество']
ROW = ['2 469,12', 'шт.', 'Кабель', '1,234.56', '2']


def font_path():
    return next(p for p in [Path('C:/Windows/Fonts/arial.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')] if p.exists())


def make_fixtures(folder):
    folder.mkdir(parents=True, exist_ok=True)
    text = '\n'.join(META) + '\n' + '\t'.join(HEAD) + '\n' + '\t'.join(ROW) + '\nИтого: 2 469,12'
    (folder / 'offer.txt').write_text(text, encoding='utf-8')
    (folder / 'offer-utf16.txt').write_text(text, encoding='utf-16')
    (folder / 'offer.csv').write_bytes((';'.join(HEAD) + '\n' + ';'.join(ROW)).encode('cp1251'))
    (folder / 'offer.tsv').write_text('\t'.join(HEAD) + '\n' + '\t'.join(ROW), encoding='utf-8-sig')
    (folder / 'offer.json').write_text(json.dumps({'client': 'ТОО Альфа', 'documentTotal': 2469.12,
        'items': [{'name': 'Кабель', 'quantity': 2, 'unit': 'шт.', 'unitPrice': 1234.56, 'lineTotal': 2469.12}]}, ensure_ascii=False), encoding='utf-8')
    (folder / 'missing.csv').write_text('Наименование;Количество;Цена;Сумма\nКабель;;;100\nБолт;2;1,234;\nИтого;;;100', encoding='utf-8')
    (folder / 'conflict.txt').write_text('Клиент: ТОО Гамма\nКлиент: ТОО Дельта\nДействительно до: 30.11.2026', encoding='utf-8')
    (folder / 'unknown.csv').write_text('SKU;Вес;Артикул\nABC;20;456', encoding='utf-8')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Товары'
    for line in META:
        ws.append([line])
    ws.append(HEAD)
    ws.append(ROW)
    ws.append(['Итого', '', '', '', '2469,12'])
    ws2 = wb.create_sheet('Услуги')
    ws2.append(['Наименование', 'Количество', 'Цена', 'Сумма'])
    ws2.append(['Монтаж', 3, 100, 300])
    wb.save(folder / 'offer.xlsx')
    legacy = xlwt.Workbook()
    sheet = legacy.add_sheet('Товары')
    for row, values in enumerate([HEAD, ROW]):
        for col, value in enumerate(values):
            sheet.write(row, col, value)
    legacy.save(str(folder / 'offer.xls'))
    doc = Document()
    for line in META:
        doc.add_paragraph(line)
    table = doc.add_table(rows=2, cols=5)
    for row, values in zip(table.rows, [HEAD, ROW]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_paragraph('Итого: 2 469,12')
    doc.save(folder / 'offer.docx')
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_font(fontname='fixture', fontfile=str(font_path()))
    for index, line in enumerate(META + ['Итого: 2 469,12']):
        page.insert_text((40, 40 + index * 20), line, fontname='fixture', fontsize=11)
    for y, values in [(350, ['Наименование', 'Кол-во', 'Ед.', 'Цена', 'Сумма']), (380, ['Кабель', '2', 'шт.', '1234,56', '2469,12'])]:
        for x, value in zip([40, 230, 290, 360, 460], values):
            page.insert_text((x, y), value, fontname='fixture', fontsize=11)
    pdf.save(folder / 'text.pdf')
    pdf.close()
    image = Image.new('RGB', (1600, 1900), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path()), 32)
    for index, line in enumerate(META + ['Итого: 2 469,12']):
        draw.text((70, 70 + index * 65), line, fill='black', font=font)
    for fmt in ['png', 'jpg', 'webp']:
        image.save(folder / f'scan.{fmt}')
    scanned = fitz.open()
    page = scanned.new_page(width=600, height=720)
    stream = BytesIO()
    image.save(stream, format='PNG')
    page.insert_image(page.rect, stream=stream.getvalue())
    scanned.save(folder / 'scan.pdf')
    scanned.close()
    with fitz.open(folder / 'text.pdf') as original:
        pix = original[0].get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False)
        pix.save(folder / 'table-scan.png')
        mixed = fitz.open()
        mixed.insert_pdf(original)
        page = mixed.new_page()
        page.insert_image(page.rect, stream=pix.tobytes('png'))
        mixed.save(folder / 'mixed.pdf')
        mixed.close()
    rotated = image.rotate(90, expand=True)
    rotated.save(folder / 'rotated.png')
    kaz = Image.new('RGB', (1500, 400), 'white')
    ImageDraw.Draw(kaz).text((40, 70), 'Жеткізуші: Қазақ Өнім\nКлиент: Әділ Ұйым', font=font, fill='black')
    kaz.save(folder / 'kazakh.png')


if __name__ == '__main__':
    make_fixtures(Path('test-results/fixtures'))
