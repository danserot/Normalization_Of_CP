import os
import re
import zipfile
from io import BytesIO
from unittest.mock import patch

import pytest
import fitz

from backend.extraction import DocumentError, read_document
from backend.pipeline import process_document
from backend.rules import extract_rules, parse_number
from backend.tests.fixtures import make_fixtures


@pytest.fixture(scope='module')
def fixtures(tmp_path_factory):
    path = tmp_path_factory.mktemp('documents')
    make_fixtures(path)
    return path


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setenv('LOCAL_MODEL_ENABLED', 'false')
    def fixture_model(source, _content, _filename):
        proposal, proof, warnings, used = extract_rules(source)
        return (proposal, proof, warnings, used), {
            'mode': 'model', 'reviewCompleted': True, 'visionUsed': False,
            'coverageComplete': True, 'unclaimedRows': [], 'issues': [],
        }
    monkeypatch.setattr('backend.pipeline.extract_universal', fixture_model)


@pytest.mark.parametrize('value,expected', [('1 234,56', 1234.56), ('1,234.56', 1234.56),
    ('1.234,56', 1234.56), ('1234.56', 1234.56), ('0', 0), ('1,234', None),
    ('29.09.2026', None), ('SKU-12', None), ('', None), ('NaN', None), ('1 234,56 KZT', 1234.56), ('12 34', None)])
def test_numbers(value, expected):
    assert parse_number(value) == expected


@pytest.mark.parametrize('name,count', [('offer.xlsx', 2), ('offer.xls', 1), ('offer.csv', 1),
    ('offer.tsv', 1), ('offer.json', 1), ('offer.docx', 1), ('offer.txt', 1), ('offer-utf16.txt', 1), ('text.pdf', 1)])
def test_formats_and_every_value_evidence(fixtures, name, count):
    result = process_document((fixtures / name).read_bytes(), name)
    p, m = result['proposal'], result['metadata']
    assert len(p['items']) == count
    assert p['items'][0] == {'name': 'Кабель', 'quantity': 2, 'unit': 'шт.', 'unitPrice': 1234.56, 'lineTotal': 2469.12}
    source = {c['id']: c for c in m['sourceCells']}
    for key, value in p.items():
        if key in ('notes', 'items') or value in ('', None):
            continue
        assert key in m['fieldEvidence'], key
    for index, item in enumerate(p['items']):
        for key, value in item.items():
            if value not in ('', None):
                assert m['fieldEvidence'][f'items.{index}.{key}']['value'] == value
    for e in m['fieldEvidence'].values():
        for entry in e if isinstance(e, list) else [e]:
            assert entry['excerpt'] == source[entry['id']]['text']
            assert entry['file'] == name
            assert entry['row'] >= 1 and entry['cell'] >= 1
    if name in ('offer.docx', 'offer.txt', 'offer.xlsx', 'text.pdf'):
        assert p['client'] == 'ТОО Альфа'
        assert p['validUntil'] == '31.10.2026'
        assert p['paymentTerms'] == '50% предоплата'
        assert p['documentTotal'] == 2469.12


def test_missing_and_ambiguous_never_default(fixtures):
    result = process_document((fixtures / 'missing.csv').read_bytes(), 'missing.csv')
    first, second = result['proposal']['items']
    assert first['quantity'] is None and first['unitPrice'] is None
    assert second['quantity'] == 2 and second['unitPrice'] is None
    assert len(result['metadata']['warnings']) >= 3


def test_arbitrary_numeric_columns_are_not_prices(fixtures):
    result = process_document((fixtures / 'unknown.csv').read_bytes(), 'unknown.csv')
    assert result['proposal']['items'] == []
    assert result['metadata']['status'] == 'empty'


def test_conflicting_values_preserved(fixtures):
    result = process_document((fixtures / 'conflict.txt').read_bytes(), 'conflict.txt')
    assert [e['value'] for e in result['metadata']['fieldEvidence']['client']] == ['ТОО Гамма', 'ТОО Дельта']


def test_invalid_signature_and_limits():
    with pytest.raises(DocumentError):
        read_document(b'not a pdf', 'bad.pdf')
    with pytest.raises(DocumentError):
        read_document(b'a' * 200001, 'huge.txt')


def test_native_text_pdf_keeps_native_values_without_vision(fixtures, monkeypatch):
    monkeypatch.setenv('VISION_ENABLED', 'false')
    result = process_document((fixtures / 'text.pdf').read_bytes(), 'text.pdf')
    assert all(cell['method'] == 'native' for cell in result['metadata']['sourceCells'])


def test_images_and_scanned_pdf_report_unreadable_when_api_unavailable(fixtures, monkeypatch):
    monkeypatch.setenv('VISION_ENABLED', 'false')
    for filename in ('scan.png', 'scan.pdf'):
        source = read_document((fixtures / filename).read_bytes(), filename)
        assert source.routing['mandatory']
        assert source.routing['available'] is False
        assert source.routing['issues']
        assert not source.text().strip()


@pytest.mark.skipif(not os.getenv('TEST_VISION'), reason='Enable TEST_VISION=1 with an OpenAI API key')
def test_real_pdf_vision(fixtures, monkeypatch):
    monkeypatch.setenv('VISION_ENABLED', 'true')
    result = process_document((fixtures / 'scan.pdf').read_bytes(), 'scan.pdf')
    assert result['metadata']['routing']['processed_pages'] == [1]
    assert result['metadata']['routing']['used']
    assert result['proposal']['client'] == 'ТОО Альфа'
    assert result['metadata']['fieldEvidence']['client']['page'] == 1
    assert result['metadata']['fieldEvidence']['client']['source_method'] == 'openai-vision'


@pytest.mark.skipif(not os.getenv('TEST_VISION'), reason='Requires a ready PaddleOCR-VL 1.6 service')
def test_visual_table_and_mixed_pdf(fixtures, monkeypatch):
    monkeypatch.setenv('VISION_ENABLED', 'true')
    with fitz.open(stream=(fixtures / 'table-scan.png').read_bytes(), filetype='png') as image:
        content = image.convert_to_pdf()
    result = process_document(content, 'table-scan.pdf')
    assert result['proposal']['items'][0]['name'] == 'Кабель'
    assert result['proposal']['items'][0]['quantity'] == 2
    assert result['proposal']['items'][0]['unitPrice'] == 1234.56
    mixed = process_document((fixtures / 'mixed.pdf').read_bytes(), 'mixed.pdf')
    assert 2 in mixed['metadata']['routing']['processed_pages']
    assert mixed['metadata']['routing']['used']
    assert len(mixed['proposal']['items']) == 2


def test_footer_is_not_item():
    result = process_document('Наименование\tКоличество\tЦена\nКабель\t2\t100\nГарантия: 12 месяцев'.encode(), 'test.txt')
    assert len(result['proposal']['items']) == 1
    assert result['proposal']['warranty'] == '12 месяцев'


def test_plain_text_is_split_into_text_and_table_blocks():
    source = read_document(
        ('Поставщик: ООО Синтетика\n'
         'Наименование\tКоличество\tЦена\n'
         'Кабель\t2\t100\n'
         'Гарантия: 12 месяцев').encode(),
        'mixed.txt',
    )
    kinds = {cell.kind for cell in source.cells}
    assert kinds == {'text', 'table'}
    assert {cell.block for cell in source.cells if cell.kind == 'text'}.isdisjoint(
        {cell.block for cell in source.cells if cell.kind == 'table'}
    )
    assert all(cell.kind == 'table' for cell in read_document(
        'Товар;Количество;Цена\nКабель;2;100'.encode(), 'table.csv').cells)


@pytest.mark.skipif(not os.getenv('TEST_VISION'), reason='Requires a ready multilingual PaddleOCR-VL 1.6 service')
def test_kazakh_characters(fixtures, monkeypatch):
    monkeypatch.setenv('VISION_ENABLED', 'true')
    with fitz.open(stream=(fixtures / 'kazakh.png').read_bytes(), filetype='png') as image:
        content = image.convert_to_pdf()
    result = process_document(content, 'kazakh.pdf')
    assert result['proposal']['client'] == 'Әділ Ұйым'
    assert result['proposal']['supplier'] == 'Қазақ Өнім'


@pytest.mark.parametrize('columns', [('name','quantity','unitPrice'), ('unitPrice','name','quantity'), ('quantity','unitPrice','name')])
def test_column_permutations(columns):
    values = {'name': 'Кабель', 'quantity': '2', 'unitPrice': '100'}
    text = ';'.join(columns) + '\n' + ';'.join(values[c] for c in columns)
    assert process_document(text.encode(), 'test.csv')['proposal']['items'][0]['unitPrice'] == 100


def test_multirow_alternative_cost_matrix_keeps_all_services_and_variants():
    from backend.extraction import Source
    source = Source('matrix.pdf')
    for row, values in enumerate([
        ['№', 'Наименование работ', 'Вариант A', '', 'Вариант B', ''],
        ['', '', 'Стоимость', 'Стоимость', 'Стоимость', 'Стоимость'],
        ['', '', 'СМР', 'ТМЦ', 'СМР', 'ТМЦ'],
        ['1', 'Монтаж блока 1', '100', '200', '90', '180'],
        ['2', 'Монтаж блока 2', '300', '400', '270', '360'],
        ['', 'Всего', '400', '', '360', ''],
    ], 1):
        for col, value in enumerate(values, 1):
            source.add(value, 'table-1', row, col, kind='table')
    proposal, proof, warnings, _ = extract_rules(source)
    assert [item['name'] for item in proposal['items']] == ['Монтаж блока 1', 'Монтаж блока 2']
    assert [component['lineTotal'] for component in proposal['items'][0]['components']] == [100, 200, 90, 180]
    assert proposal['documentTotal'] is None
    assert [field['value'] for field in proposal['additionalFields']] == ['400', '360']
    assert all(f'items.{index}.name' in proof for index in range(2))
    assert any('альтернативные варианты' in warning for warning in warnings)


@pytest.mark.parametrize('dimension', ['', '<dimension ref="A1:A1"/>'])
def test_xlsx_without_reliable_dimensions(fixtures, dimension):
    output = BytesIO()
    with zipfile.ZipFile(fixtures / 'offer.xlsx') as original, zipfile.ZipFile(output, 'w') as rewritten:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename.startswith('xl/worksheets/sheet'):
                data = re.sub(rb'<dimension\b[^>]*/>', dimension.encode(), data)
            rewritten.writestr(entry, data)
    result = process_document(output.getvalue(), 'no-dimensions.xlsx')
    assert len(result['proposal']['items']) == 2
    assert result['proposal']['items'][0]['unitPrice'] == 1234.56
    assert result['metadata']['fieldEvidence']['items.1.name']['sheet'] == 'Услуги'


@pytest.mark.parametrize('row,col', [(10001, 1), (1, 101)])
def test_xlsx_unknown_dimensions_still_enforces_limits(row, col):
    from openpyxl import Workbook
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    for _ in range(row - 1):
        sheet.append([])
    sheet.append([None] * (col - 1) + ['data'])
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(DocumentError, match='10 000 строк / 100 колонок'):
        read_document(output.getvalue(), 'limits.xlsx')


def test_xlsx_footer_after_eight_items_is_not_ninth_item():
    from openpyxl import Workbook
    workbook = Workbook(write_only=True)  # Valid XLSX without a dimension element.
    sheet = workbook.create_sheet()
    sheet.append(['№', 'Наименование', 'Кол-во', 'Ед. изм.', 'Цена за ед., ₸', 'Сумма, ₸'])
    for index in range(8):
        sheet.append([index + 1, f'Товар {index + 1}', 2, 'шт.', 100, 200])
    sheet.append(['ИТОГО:', None, None, None, None, 1600])
    sheet.append(['Контактное лицо', 'Тестовый контакт'])
    sheet.append(['Отдел продаж', 'Служебные сведения'])
    output = BytesIO()
    workbook.save(output)
    result = process_document(output.getvalue(), 'eight-items.xlsx')
    assert len(result['proposal']['items']) == 8
    assert result['proposal']['documentTotal'] == 1600
    assert result['proposal']['clientContact'] == 'Тестовый контакт'
