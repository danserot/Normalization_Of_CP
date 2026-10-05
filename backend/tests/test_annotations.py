"""Original previews, durable review, source grounding and grouped SFT export."""
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient

from backend import main
from backend.annotation_data import Annotation, prepare_annotation, training_examples, validate_annotation
from backend.extraction import Source
from backend.universal import Layout, Requisites, metadata_cells
from backend.universal import catalog
from backend.annotation_data import word_preview

SAMPLES = Path(__file__).resolve().parents[2] / 'examples' / 'mock_kp'


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(main, 'DATABASE_PATH', tmp_path / 'annotations.sqlite3')
    monkeypatch.setattr(main, 'APP_PASSWORD', '')
    monkeypatch.setenv('LOCAL_MODEL_ENABLED', 'false')
    main.initialize_db()
    with TestClient(main.app) as instance:
        yield instance


def upload(client, filename='offer.txt', content=None):
    content = content if content is not None else (SAMPLES / filename).read_bytes()
    response = client.post('/api/annotations', files={'file': (filename, content)})
    assert response.status_code == 200, response.text
    return response.json()


def mark_reviewed(document, group='same-template'):
    annotation = document['annotation']
    for field in annotation['fields']:
        field.update(state='missing', cell='', value='')
    supplier = next(c for c in document['cells'] if 'Поставщик:' in c['text'])
    next(f for f in annotation['fields'] if f['field'] == 'supplier').update(
        state='found', cell=supplier['id'], value=supplier['text'].split(':', 1)[1].strip())
    for table in annotation['tables']:
        table.update(reviewed=True, isItems=False)
    annotation['group'] = group
    return {'revision': document['revision'], 'status': 'reviewed', 'annotation': annotation}


def test_original_deduplication_persistence_and_revision_conflict(client):
    original = (SAMPLES / 'offer.txt').read_bytes()
    first = upload(client)
    again = upload(client)
    assert again['duplicate'] and first['id'] == again['id']
    identifier = first['id']
    assert client.get(f'/api/annotations/{identifier}/original').content == original
    assert client.get('/api/annotations').json()['documents'][0]['id'] == identifier
    submission = mark_reviewed(first)
    saved = client.post(f'/api/annotations/{identifier}', json=submission)
    assert saved.status_code == 200, saved.text
    assert saved.json()['revision'] == 2
    main.initialize_db()
    assert client.get(f'/api/annotations/{identifier}').json()['annotation'] == submission['annotation']
    assert client.post(f'/api/annotations/{identifier}', json=submission).status_code == 409


def test_review_rejects_unchecked_or_invented_answers_but_draft_is_allowed(client):
    document = upload(client)
    identifier = document['id']
    pending = {'revision': 1, 'status': 'reviewed', 'annotation': document['annotation']}
    assert client.post(f'/api/annotations/{identifier}', json=pending).status_code == 422
    submission = mark_reviewed(document)
    supplier = next(f for f in submission['annotation']['fields'] if f['field'] == 'supplier')
    supplier['value'] = 'Несуществующая компания'
    assert client.post(f'/api/annotations/{identifier}', json=submission).status_code == 422
    submission['status'] = 'draft'
    assert client.post(f'/api/annotations/{identifier}', json=submission).status_code == 200
    assert client.get('/api/annotations/export').status_code == 422


def test_export_only_reviewed_documents_with_groups_and_production_schemas(client):
    first = upload(client)
    second = upload(client, 'other.txt', ('Поставщик: ООО Вторая\nКлиент: ООО Третья').encode())
    draft = upload(client, 'draft.txt', b'Client: Not Reviewed')
    for document in (first, second):
        response = client.post('/api/annotations/' + document['id'], json=mark_reviewed(document))
        assert response.status_code == 200, response.text
    exported = client.get('/api/annotations/export')
    assert exported.status_code == 200
    with ZipFile(BytesIO(exported.content)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        assert len(manifest['documents']) == 2
        assert draft['id'] not in {d['id'] for d in manifest['documents']}
        assert len({d['split'] for d in manifest['documents']}) == 1
        assert sum(manifest['counts'].values()) > 0
        for split in ('train', 'val', 'test'):
            for example in json.loads(archive.read(f'kp_{split}.json')):
                answer = json.loads(example['output'])
                (Requisites if 'fields' in answer else Layout).model_validate(answer)
                assert json.loads(example['instruction'])
                assert example['system']
        assert 'kp_test' in json.loads(archive.read('dataset_info.json'))


@pytest.mark.parametrize('filename,kind', [
    ('offer.pdf', 'pdf'), ('offer-scan.pdf', 'pdf'), ('offer.docx', 'word'),
    ('offer.xlsx', 'table'), ('offer.xls', 'table'), ('offer.csv', 'table'), ('offer.tsv', 'table'),
    ('offer.txt', 'text'), ('offer.json', 'text'), ('offer.png', 'image'), ('offer.jpg', 'image'), ('offer.webp', 'image'),
])
def test_all_supported_format_previews(client, filename, kind):
    document = upload(client, filename)
    assert document['preview']['kind'] == kind
    assert document['cells'], document['warnings']
    if kind in ('pdf', 'image'):
        preview = client.get(f"/api/annotations/{document['id']}/pages/1")
        assert preview.status_code == 200
        assert preview.headers['content-type'].startswith('image/')
        assert client.get(f"/api/annotations/{document['id']}/pages/999").status_code == 404
    if kind == 'word':
        preview = client.get(f"/api/annotations/{document['id']}/preview")
        assert preview.status_code == 200 and '<table>' in preview.text
        assert "default-src 'none'" in preview.headers['content-security-policy']


def test_requisites_above_and_below_items_remain_trainable():
    payload = prepare_annotation((SAMPLES / 'offer.xlsx').read_bytes(), 'offer.xlsx')
    annotation = Annotation.model_validate(payload.pop('annotation'))
    for field in annotation.fields:
        field.state = 'found' if field.cell and field.value else 'missing'
    annotation.group = 'template'
    for table in annotation.tables:
        table.reviewed = True
        # The final row is a document total, not a product.
        table.lastRow -= 1
    validate_annotation(annotation, payload, 'offer.xlsx', True)
    examples = training_examples(annotation, payload, 'offer.xlsx')
    found = [field for example in examples if 'fields' in json.loads(example['output'])
             for field in json.loads(example['output'])['fields']]
    assert any(f['field'] == 'supplier' for f in found)
    assert any(f['field'] == 'documentTotal' for f in found)


def test_annotation_routes_require_login(client, monkeypatch):
    monkeypatch.setattr(main, 'APP_PASSWORD', 'secret')
    assert client.get('/api/annotations').status_code == 401
    assert client.post('/api/annotations', files={'file': ('x.txt', b'text')}).status_code == 401
    assert client.get('/api/annotations/export').status_code == 401


def test_requisites_filter_excludes_only_product_rows():
    source = Source('sample.xlsx')
    source.add('Поставщик: Альфа', 'sheet', 1)
    source.add('Товар', 'sheet', 2, 1)
    source.add('100', 'sheet', 2, 2)
    source.add('Итого: 100', 'sheet', 3)
    table = type('Table', (), {'block': 'sheet', 'firstRow': 2, 'lastRow': 2})()
    assert [c.row for c in metadata_cells(source, [table])] == [1, 3]


def test_late_table_header_is_visible_to_model_and_training():
    source = Source('long.xlsx')
    for row in range(1, 26):
        source.add('Реквизит', 'sheet', row)
    for col, text in enumerate(['Товар', 'Количество', 'Цена'], 1):
        source.add(text, 'sheet', 26, col)
    for col, text in enumerate(['Кабель', '2', '100'], 1):
        source.add(text, 'sheet', 27, col)
    source.add('Итог', 'sheet', 28)
    assert 26 in {row['row'] for row in catalog(source)[0]['rows']}


def test_word_preview_preserves_images_hyperlink_text_nested_tables_and_escapes_markup():
    from docx import Document
    from docx.oxml import OxmlElement
    from PIL import Image
    document = Document()
    paragraph = document.add_paragraph('<script>unsafe</script>')
    link = OxmlElement('w:hyperlink')
    run, text = OxmlElement('w:r'), OxmlElement('w:t')
    text.text = 'Текст ссылки'
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    image = BytesIO()
    Image.new('RGB', (30, 30), 'white').save(image, format='PNG')
    image.seek(0)
    document.add_picture(image)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = 'Вложенная таблица'
    output = BytesIO()
    document.save(output)
    html = word_preview(output.getvalue())
    assert '<script>' not in html and '&lt;script&gt;' in html
    assert 'Текст ссылки' in html and 'Вложенная таблица' in html
    assert 'data:image/jpeg;base64,' in html
